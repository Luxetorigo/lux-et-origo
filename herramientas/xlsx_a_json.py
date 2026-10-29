#!/usr/bin/env python3
"""Convierte el Excel de Lux et Origo en los JSON que lee el sitio.

Uso:  python3 herramientas/xlsx_a_json.py MiExcel.xlsx
Lee las hojas «Partituras», «Recursos» y «Salmos» (otras hojas se ignoran) y escribe
data/partituras.json, data/recursos.json y data/salmos.json.
Regla del Excel: un campo vacío se escribe 0 (o se deja en blanco). El sitio decide
cómo mostrarlo («No disponible», «Anónimo», etc.).
"""
import json, re, sys, unicodedata
from pathlib import Path
from openpyxl import load_workbook

def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).lower()
    return re.sub(r"[^a-z0-9]+", " ", "".join(c for c in s if unicodedata.category(c) != "Mn")).strip()

def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", norm(s)).strip("-")

def val(c):
    """Celda -> texto limpio. None, vacío y 0 significan «sin dato»."""
    if c is None: return ""
    if isinstance(c, float) and c.is_integer(): c = int(c)
    s = str(c).strip()
    return "" if s in ("", "0", "0.0") else s

HOJAS = {
    "partituras": ("data/partituras.json", {
        "titulo": ["titulo", "nombre"], "tipo": ["tipo"], "autor": ["autor", "compositor"],
        "parte": ["parte", "parte de la misa"],
        "drive_reducida": ["drive reducida", "drive partitura"],
        "drive_original": ["drive original"],
        "youtube": ["youtube"], "fuente": ["fuente"], "condicion": ["condicion"],
        "autorizada_por": ["autorizada por"], "tiempo": ["tiempo", "tiempo liturgico"],
        "descripcion": ["descripcion"], "id": ["id"],
    }),
    "recursos": ("data/recursos.json", {
        "categoria": ["categoria", "tipo"], "titulo": ["titulo", "nombre"], "autor": ["autor"],
        "descripcion": ["descripcion"], "drive": ["drive"], "fuente": ["fuente"],
        "editorial": ["editorial"], "id": ["id"],
    }),
    "salmos": ("data/salmos.json", {
        "titulo": ["titulo", "nombre"], "numero": ["no de salmo", "numero", "n de salmo", "salmo"],
        "autor": ["autor"], "rito": ["dia o rito", "rito", "dia"], "descripcion": ["descripcion"],
        "drive": ["drive"], "youtube": ["youtube"], "fuente": ["fuente"], "id": ["id"],
    }),
}

BASE = "https://luxetorigo.github.io/lux-et-origo/"

def generar_sitemap(raiz):
    """sitemap.xml con las páginas fijas y la página de cada partitura, recurso y salmo."""
    from datetime import date
    hoy = date.today().isoformat()
    urls = ["", "recursos.html", "salmos.html", "transparencia.html"]
    for archivo, pagina in (("partituras", "partitura.html"), ("recursos", "recurso.html"), ("salmos", "salmo.html")):
        try:
            for it in json.loads((raiz / "data" / f"{archivo}.json").read_text(encoding="utf-8")):
                urls.append(f"{pagina}?id={it['id']}")
        except FileNotFoundError:
            pass
    cuerpo = "".join(f"  <url><loc>{BASE}{u}</loc><lastmod>{hoy}</lastmod></url>\n" for u in urls)
    (raiz / "sitemap.xml").write_text('<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n' + cuerpo + "</urlset>\n", encoding="utf-8")
    (raiz / "robots.txt").write_text(f"User-agent: *\nAllow: /\nSitemap: {BASE}sitemap.xml\n", encoding="utf-8")
    print(f"sitemap.xml: {len(urls)} dirección(es)")

def categoria(txt):
    n = norm(txt)
    for pref, nombre in (("cantoral", "Cantorales"), ("libro", "Libros"), ("documento", "Documentos")):
        if n.startswith(pref): return nombre
    return ""

def convertir(xlsx, raiz):
    wb = load_workbook(xlsx, data_only=True)
    hojas = {norm(n): wb[n] for n in wb.sheetnames}
    avisos = []
    for clave, (destino, campos) in HOJAS.items():
        out, ids = [], set()
        ws = hojas.get(clave)
        if ws is None:
            avisos.append(f"No encontré la hoja «{clave.capitalize()}» (se deja vacía).")
        else:
            filas = list(ws.iter_rows(values_only=True))
            cab = [norm(c) for c in (filas[0] if filas else [])]
            idx = {}
            for campo, alias in campos.items():
                for a in alias:
                    if a in cab: idx[campo] = cab.index(a); break
            for n, fila in enumerate(filas[1:], start=2):
                reg = {c: (val(fila[i]) if i < len(fila) else "") for c, i in idx.items()}
                if not any(reg.values()): continue
                t = reg.get("titulo", "")
                if not t: avisos.append(f"{clave} fila {n}: sin título, no aparece."); continue
                if t.upper().startswith("EJEMPLO"): continue
                if clave == "recursos":
                    bruto = reg.get("categoria", "")
                    reg["categoria"] = categoria(bruto)
                    if bruto and not reg["categoria"]:
                        avisos.append(f"recursos fila {n} («{t}»): categoría «{bruto}» no válida (Cantorales, Libros o Documentos); queda sin clasificar.")
                base = reg.get("id") or slug(t)
                i, nid = 2, base
                while nid in ids: nid = f"{base}-{i}"; i += 1
                ids.add(nid); reg["id"] = nid
                for k in ("drive", "drive_reducida", "drive_original"):
                    v = reg.get(k, "")
                    if v and "drive.google.com" not in v and "docs.google.com" not in v:
                        avisos.append(f"{clave} fila {n} («{t}»): el enlace de {k} no parece de Drive.")
                v = reg.get("youtube", "")
                if v and "youtu" not in v: avisos.append(f"{clave} fila {n} («{t}»): el enlace de YouTube no parece válido.")
                out.append({"id": reg.pop("id"), **reg})
        (raiz / destino).write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"{destino}: {len(out)} elemento(s)")
    generar_sitemap(raiz)
    for a in avisos: print("AVISO:", a)

if __name__ == "__main__":
    if len(sys.argv) != 2: sys.exit(__doc__)
    convertir(sys.argv[1], Path(__file__).resolve().parent.parent)
